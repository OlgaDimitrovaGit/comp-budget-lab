# -*- coding: utf-8 -*-
"""
Builds merit-model.xlsx from a CSV of employees.

Every figure is a live formula. Calc runs left to right in numbered steps, one
simple operation per column; rows 1-4 carry the step, a note, an optional
second-language note and the column header.
"""
import csv
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

from reference_calc import DEFAULTS, load_bands, parse_date

SRC = sys.argv[1] if len(sys.argv) > 1 else "demo-data.csv"
BANDS = sys.argv[2] if len(sys.argv) > 2 else "demo-bands.csv"
OUT = sys.argv[3] if len(sys.argv) > 3 else "merit-model.xlsx"

INPUT = PatternFill("solid", fgColor="FFF4CC")
AMBER = PatternFill("solid", start_color="FFC000", end_color="FFC000")
STEP = PatternFill("solid", fgColor="DCE6F0")
BOLD = Font(bold=True)
RU = Font(italic=True, color="7F7F7F")
WRAP = Alignment(wrap_text=True, vertical="top")
EUR = '#,##0.00 "€"'
PCT = "0.00%"
DATE = "yyyy-mm-dd"
MONTH = "mmm yyyy"

wb = Workbook()


def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def note2(ws, r, c, text, font=RU, wrap=True):
    """Second-language note: the cell is written only when there is a note."""
    if text:
        cell = ws.cell(r, c, text)
        cell.font = font
        if wrap:
            cell.alignment = WRAP


# ---------------------------------------------------------------- Inputs
ws = wb.active
ws.title = "Inputs"
ws["A1"] = "Inputs: change the yellow cells"
ws["A1"].font = Font(bold=True, size=13)
ws.append([])
ws.append(["Setting", "Value", "Note"])
note2(ws, 3, 4, "", wrap=False)
for c in ws[3]:
    c.font = BOLD

S = DEFAULTS
settings = [
    ("EffDate", "Effective from (pick a month)", S["effective_date"], MONTH, True,
     "Increases take effect on the 1st of this month. Eligibility is counted up to it. Pick from the list.",
     ""),
    ("EffMonth", "Effective month", "=MONTH(EffDate)", "0", False,
     "Month number of the effective date.", ""),
    ("InYearFactor", "In-year proration factor", "=(13-EffMonth)/12", PCT, False,
     "Months from the effective month to December inclusive, ÷ 12; the increase is assumed to start on the 1st of that month.",
     ""),
    ("PoolPct", "Target merit increase, % of eligible base payroll", S["pool_pct"] / 100, PCT, True,
     "The planned average increase, as % of eligible staff's annual base after the pay equity step, FTE-weighted. "
     "The target budget on Summary is this % given to every eligible person, with bonus and employer contributions.",
     ""
     ""),
    ("MinMonthsHire", "Minimum tenure, months", S["min_months_hire"], '0 "mo"', True,
     "Hired fewer months before the effective date than this: not in the review. Hired after the effective date: never in it.",
     ""),
    ("MinMonthsChange", "No individual pay change in the last, months", S["min_months_change"], '0 "mo"', True,
     "Individual pay change fewer months before the effective date than this: not in the review.",
     ""),
    ("MaxIncrease", "Increase cap per person, % of aligned base", S["max_increase"] or "No cap", PCT, True,
     "A safety limit. The increase is each person's gap to target, or the rating minimum if higher; the cap only cuts it if it is larger. "
     "'No cap' (or an empty cell, or 0%) = no limit.",
     ""
     ""),
    ("ClipAtMax", "Clip increases at range maximum", "Yes" if S["clip_at_max"] else "No", "@", True,
     "Yes: a merit increase never takes base pay above the range maximum. The pay equity adjustment is not clipped.",
     ""),
    ("CompanyTarget", "Pay policy position, % of range midpoint", S["company_target"], PCT, True,
     "Leave at 100% if your range midpoints already reflect your pay policy.",
     ""),
    ("RateBelow", "Employer contribution rate up to the ceiling", S["rate_below"] / 100, PCT, True,
     "Default: Spain 2026, 30.65% (permanent contracts; fixed-term 31.85%) + AT/EP occupational accident and disease premium, office-work rate 1.50% (RDL 3/2026, Cuadro II a) = 32.15%. Use your activity's AT/EP rate. These are 2026 rates. Rates for 2027 are already set and slightly higher (MEI employer share 0.83%, higher solidarity rates); 2026 rates are kept, so the contribution cost of a 2027 review is slightly understated.",
     ""),
    ("Ceiling", "Annual contribution ceiling", S["ceiling"], EUR, True,
     "Default: Spain 2026, €5,101.20 a month × 12 (Orden PJC/297/2026). 2027 not yet published: 2026 assumed.",
     ""),
    ("AboveCeilingTier1", "Rate above the ceiling, tier 1: up to 10% over the ceiling", S["above_ceiling"][0][1] / 100, PCT, True,
     "Employer rate on the part of pay above the ceiling. Set 0% in all three tiers if your country "
     "charges nothing above it. Default: Spain 2026 (cuota de solidaridad), employer share, 0.96%.",
     ""
     ""),
    ("AboveCeilingTier2", "Rate above the ceiling, tier 2: 10% to 50% over the ceiling", S["above_ceiling"][1][1] / 100, PCT, True,
     "Default: Spain 2026, employer share, 1.04%.",
     ""),
    ("AboveCeilingTier3", "Rate above the ceiling, tier 3: more than 50% over the ceiling", S["above_ceiling"][2][1] / 100, PCT, True,
     "Default: Spain 2026, employer share, 1.22%.",
     ""),
]
assert [t for t, _ in S["above_ceiling"]] == [1.10, 1.50, None]  # tier bounds are fixed in the formula
for i, (n, label, val, fmt, is_input, en, ru) in enumerate(settings):
    r = 4 + i
    ws.cell(r, 1, label)
    c = ws.cell(r, 2, val)
    c.number_format = fmt
    if is_input:
        c.fill = INPUT
    ws.cell(r, 3, en).alignment = WRAP
    note2(ws, r, 4, ru)
    name(n, f"Inputs!$B${r}")

RT = 4 + len(settings) + 2  # rating table header row
ws.cell(RT - 1, 1, "Targets by performance rating (illustrative assumption: set your own)").font = BOLD
note2(ws, RT - 1, 4, "", wrap=False)
for j, h in enumerate(["Rating", "Label", "Target compa-ratio", "Minimum increase, % of aligned base"]):
    ws.cell(RT, 1 + j, h).font = BOLD
labels = {1: "Unsatisfactory", 2: "Partially meets", 3: "Meets", 4: "Exceeds", 5: "Outstanding"}
for k in range(1, 6):
    r = RT + k
    tgt, mn = S["ratings"][k]
    ws.cell(r, 1, k)
    ws.cell(r, 2, labels[k])
    for j, v in enumerate((tgt, mn)):
        c = ws.cell(r, 3 + j, v)
        c.number_format = PCT
        c.fill = INPUT
R1, R2 = RT + 1, RT + 5
notes = [
    ("Target compa-ratio: the compa-ratio each rating is raised to. The increase is the gap "
     "from aligned base to mid × pay policy position × this target. Set 0% to give a rating no increase at all. "
     "A target above a grade's range max (max / mid) is cut at the max when 'Clip increases at range maximum' = Yes.",
     ""
     ""
     ""),
    ("Minimum increase: a floor for a rating, paid even when the person is already at or above "
     "the target. 0% = no floor. It does not apply at or above the range maximum or to ineligible staff, "
     "and it is limited by the per-person cap and, when the clip at max is on, by the headroom to the max.",
     ""
     ""
     ""),
]
for i, (en, ru) in enumerate(notes):
    ws.cell(R2 + 1 + i, 1, en).alignment = WRAP
    note2(ws, R2 + 1 + i, 4, ru)

ZT = R2 + 5  # zone table header row
ws.cell(ZT - 1, 1, "Compa-ratio zones: only for the Matrix sheet view").font = BOLD
note2(ws, ZT - 1, 4, ""
      "")
ws.cell(ZT - 1, 3, "Increases do not depend on zones. Zones only group people into the classic "
        "rating × compa-ratio grid on the Matrix sheet.").alignment = WRAP
for j, h in enumerate(["Zone", "Lower bound (CR from)"]):
    ws.cell(ZT, 1 + j, h).font = BOLD
for k, z in enumerate(S["zones"]):
    ws.cell(ZT + 1 + k, 1, k + 1)
    c = ws.cell(ZT + 1 + k, 2, z)
    c.number_format = PCT
    c.fill = INPUT
Z1, Z2 = ZT + 1, ZT + len(S["zones"])
for col, w in zip("ABCDE", (46, 16, 60, 60, 22)):
    ws.column_dimensions[col].width = w

# ---------------------------------------------------------------- Lists (drop-down values for Inputs)
wl = wb.create_sheet("Lists")
LISTS = {
    "EffDate": ("Effective from", [date(2026 + (10 + i) // 12, (10 + i) % 12 + 1, 1) for i in range(26)], MONTH),
    "MinMonthsHire": ("Minimum tenure, months", [0, 1, 2, 3, 6, 9, 12], '0 "mo"'),
    "MinMonthsChange": ("No pay change in the last, months", [0, 1, 2, 3, 6, 9, 12], '0 "mo"'),
    "MaxIncrease": ("Increase cap", ["No cap", 0.10, 0.15, 0.20, 0.25, 0.30, 0.40, 0.50, 0.75, 1.00], PCT),
    "ClipAtMax": ("Clip at range max", ["Yes", "No"], "@"),
}
setting_row = {n: 4 + i for i, (n, *_) in enumerate(settings)}
for j, (key, (head, values, fmt)) in enumerate(LISTS.items()):
    col = get_column_letter(1 + j)
    wl.cell(1, 1 + j, head).font = BOLD
    for i, v in enumerate(values):
        wl.cell(2 + i, 1 + j, v).number_format = fmt
    wl.column_dimensions[col].width = 18
    dv = DataValidation(type="list", formula1=f"=Lists!${col}$2:${col}${1 + len(values)}", allow_blank=True,
                        showErrorMessage=True, errorTitle="Pick from the list",
                        error="Choose a value from the drop-down list.")
    ws.add_data_validation(dv)
    dv.add(f"B{setting_row[key]}")
assert S["effective_date"] in LISTS["EffDate"][1]
wl.sheet_state = "hidden"

# ---------------------------------------------------------------- Bands
wb_b = wb.create_sheet("Bands")
wb_b.append(["Category", "Grade", "Range min", "Range mid", "Range max", "Lookup key"])
for c in wb_b[1]:
    c.font = BOLD
for i, ((cat, g), (mn, md, mx)) in enumerate(load_bands(BANDS).items()):
    r = 2 + i
    wb_b.append([cat, g, mn, md, mx, f'=A{r}&"|"&B{r}'])
for row in wb_b.iter_rows(min_row=2, max_col=5):
    for c in row:
        c.fill = INPUT
    for c in row[2:]:
        c.number_format = EUR
wb_b.column_dimensions["A"].width = 22
for col in "BCDEF":
    wb_b.column_dimensions[col].width = 14

# ---------------------------------------------------------------- Employees
we = wb.create_sheet("Employees")
cols = ["id", "category", "gender", "grade", "rating", "hire_date", "last_change_date",
        "base_salary", "bonus_pct", "fte"]
we.append(cols)
for c in we[1]:
    c.font = BOLD
rows = list(csv.DictReader(open(SRC, encoding="utf-8-sig")))
for r in rows:
    we.append([r["id"], r["category"], r["gender"], int(r["grade"]), int(r["rating"]),
               parse_date(r["hire_date"]), parse_date(r["last_change_date"]),
               float(r["base_salary"]), float(r["bonus_pct"]), float(r["fte"])])
for row in we.iter_rows(min_row=2):
    for c in row:
        c.fill = INPUT
    row[5].number_format = row[6].number_format = DATE
    row[7].number_format = EUR
for col, w in zip("ABCDEFGHIJ", (8, 22, 8, 7, 7, 12, 16, 13, 10, 6)):
    we.column_dimensions[col].width = w
N = len(rows)

# ---------------------------------------------------------------- Calc
def EMP(field):
    """Look a field up on Employees by ID and by column header, not by cell."""
    return f'=INDEX(Employees!$A:$Z,MATCH({{id}},Employees!$A:$A,0),MATCH("{field}",Employees!$1:$1,0))'


# Employer contributions on annual pay {x}: regressive Spanish scale
CONTRIB = ("=RateBelow*MIN({x},Ceiling)+AboveCeilingTier1*MIN(MAX(0,{x}-Ceiling),0.1*Ceiling)"
           "+AboveCeilingTier2*MIN(MAX(0,{x}-1.1*Ceiling),0.4*Ceiling)+AboveCeilingTier3*MAX(0,{x}-1.5*Ceiling)")

# (key, header, number format, formula template, note, second-language note)
# {k} -> this row's cell of column k; {er} -> Employees row
CALC = [
    ("Step 0: Input", [
        ("n", "No.", "0", "=ROW()-4", "Row number in the list.", ""),
        ("id", "ID", None, "=INDEX(Employees!$A:$A,{n}+1)", "The n-th ID on Employees.", ""),
        ("cat", "Category", None, EMP("category"), "Category of workers (same work or work of equal value); looked up by ID.", ""),
        ("gen", "Gender", None, EMP("gender"), "F or M.", ""),
        ("grade", "Grade", "0", EMP("grade"), "", ""),
        ("rating", "Rating", "0", EMP("rating"), "1 to 5.", ""),
        ("hire", "Hire date", DATE, EMP("hire_date"), "", ""),
        ("chg", "Last pay change", DATE, EMP("last_change_date"), "Last individual change.", ""),
        ("base", "Base salary", EUR, EMP("base_salary"), "Annual gross base, full time; in Spain including the extra payments (pagas extra).",
         ""),
        ("bonus", "Bonus % of base", PCT, EMP("bonus_pct") + "/100", "Bonus as a share of base.", ""),
        ("fte", "FTE", "0.00", EMP("fte"), "1 = full time.", ""),
    ]),
    ("Step 1: Range and total cash", [
        ("bok", "Range found", "0", '=IF(ISNA(MATCH({cat}&"|"&{grade},Bands!$F:$F,0)),0,1)',
         "0: no row for this category and grade on Bands; totals show #N/A until it is added.",
         ""),
        ("bmin", "Range min", EUR, '=INDEX(Bands!$C:$C,MATCH({cat}&"|"&{grade},Bands!$F:$F,0))', "From Bands by category and grade.", ""),
        ("bmid", "Range mid", EUR, '=INDEX(Bands!$D:$D,MATCH({cat}&"|"&{grade},Bands!$F:$F,0))', "", ""),
        ("bmax", "Range max", EUR, '=INDEX(Bands!$E:$E,MATCH({cat}&"|"&{grade},Bands!$F:$F,0))', "", ""),
        ("ten", "Tenure, years", "0.00", "=(EffDate-{hire})/365.25", "Years from hire to effective date.", ""),
        ("tot", "Total cash", EUR, "={base}*(1+{bonus})", "Base plus bonus, full time.", ""),
    ]),
    ("Step 2: Pay equity adjustment (unexplained gap closed in full)", [
        ("g2", "grade²", "0.00", "={grade}^2", "Helper sums for the regression on Categories.", ""),
        ("t2", "tenure²", "0.00", "={ten}^2", "", ""),
        ("gt", "grade×tenure", "0.00", "={grade}*{ten}", "", ""),
        ("gy", "grade×cash", "0.00", "={grade}*{tot}", "", ""),
        ("ty", "tenure×cash", "0.00", "={ten}*{tot}", "", ""),
        ("totr", "Total cash, rounded", EUR, "=ROUND({tot},2)", "Rounded to the cent so that ranking compares exact values.", ""),
        ("rank", "Rank in category", "0", '=COUNTIFS({C_cat},{cat},{C_totr},"<"&{totr})+COUNTIFS({C_cat},{cat},{C_totr},{totr},{C_id},"<"&{id})+1',
         "Order by total cash within the category, for the median.", ""),
        ("cmed", "Category median total cash", EUR, "=INDEX(Categories!$AH:$AH,MATCH({cat},Categories!$A:$A,0))", "From Categories.", ""),
        ("rec", "Receives adjustment", "0", '=IF(AND({gen}="F",{tot}<{cmed},INDEX(Categories!$AI:$AI,MATCH({cat},Categories!$A:$A,0))>0),1,0)',
         "1: woman below the category median, and the category has an unexplained gap against women.", ""),
        ("def", "Shortfall to median", EUR, "=IF({rec}=1,{cmed}-{tot},0)", "", ""),
        ("upt", "Adjustment, total cash", EUR, "=IF({rec}=1,INDEX(Categories!$AI:$AI,MATCH({cat},Categories!$A:$A,0))*{def}/INDEX(Categories!$AJ:$AJ,MATCH({cat},Categories!$A:$A,0)),0)",
         "Category amount shared in proportion to each shortfall.", ""),
        ("upb", "Adjustment, base", EUR, "={upt}/(1+{bonus})", "Paid through base; the bonus follows the base.", ""),
        ("al", "Aligned base", EUR, "={base}+{upb}", "Merit is computed from this base.", ""),
        ("ovmax1", "Above range max after equity adjustment", "0", "=IF(AND({upb}>0,{al}>{bmax}),1,0)",
         "Amber flag, not an error. The adjustment is paid in full. The category's amount is shared in proportion to each woman's shortfall to the category median, and a category spans several grades, so her new pay can sit above her grade's range maximum. Worth a look: her grade, the range, or how the category is defined.",
         ""),
    ]),
    ("Step 3: Eligibility", [
        ("mh", "Months since hire", "0", '=IF({hire}>EffDate,-1,DATEDIF({hire},EffDate,"m"))',
         "Complete months up to the effective date; −1 if hired after it.", ""),
        ("mc", "Months since change", "0", '=IF({chg}>EffDate,-1,DATEDIF({chg},EffDate,"m"))',
         "−1 if the change is after the effective date.", ""),
        ("elig", "Eligible", "0", "=IF(AND({mh}>=MinMonthsHire,{mc}>=MinMonthsChange),1,0)", "1: in the review.", ""),
    ]),
    ("Step 4: Target salary", [
        ("cr", "Compa-ratio", PCT, "={al}/{bmid}", "Aligned base / range mid.", ""),
        ("zone", "Compa-ratio zone", "0", f"=MATCH({{cr}},Inputs!$B${Z1}:$B${Z2},1)",
         "Not used in the increase. Only places the person in a column of the Matrix sheet grid (rating × compa-ratio).",
         ""),
        ("tcr", "Target compa-ratio for rating", PCT, f"=INDEX(Inputs!$C${R1}:$C${R2},MATCH({{rating}},Inputs!$A${R1}:$A${R2},0))", "From the rating table.", ""),
        ("minp", "Minimum increase", PCT, f"=INDEX(Inputs!$D${R1}:$D${R2},MATCH({{rating}},Inputs!$A${R1}:$A${R2},0))",
         "Floor for the rating, paid even at or above target; see Inputs.",
         ""),
        ("ts", "Target salary", EUR, "={bmid}*CompanyTarget*{tcr}", "Mid × pay policy position × target compa-ratio.", ""),
        ("gap", "Gap to target", EUR, "=MAX(0,{ts}-{al})", "0 if already at or above target.", ""),
    ]),
    ("Step 5: Merit increase", [
        ("capon", "Cap applies", "0", "=IF(AND(ISNUMBER(MaxIncrease),MaxIncrease>0),1,0)", "0: 'No cap', empty or 0% on Inputs.",
         ""),
        ("capa", "Increase cap, €", EUR, "=IF({capon}=1,{al}*MaxIncrease,0)", "Largest increase for this person: the cap × their aligned base; 0 when there is no cap.",
         ""),
        ("mina", "Minimum increase, €", EUR, "={al}*{minp}", "", ""),
        ("rule", "Increase before range limits", EUR, "=IF({tcr}=0,0,IF({capon}=1,MIN(MAX({gap},{mina}),{capa}),MAX({gap},{mina})))",
         "Gap to target, at least the minimum, never above the cap (if there is one).",
         ""),
        ("red", "At or above range max (red-circled)", "0", "=IF({al}>={bmax},1,0)", "At or above range max: no merit increase.", ""),
        ("afr", "After red-circle rule", EUR, "=IF({red}=1,0,{rule})", "Already at or above max: increase set to 0, whatever the rule gave.",
         ""),
        ("room", "Headroom to range max", EUR, "=MAX(0,{bmax}-{al})", "Range max − aligned base: the most an increase can add without leaving the range.",
         ""),
        ("afc", "After clip at max", EUR, "=IF(ClipAtMax=\"No\",{afr},MIN({afr},{room}))", "If 'Clip increases at range maximum' = Yes, an increase larger than the headroom is cut to the headroom, so new base = range max.",
         ""),
        ("inc", "Merit increase", EUR, "=IF({elig}=1,{afc},0)", "Annual full-time base increase.", ""),
        ("incp", "Increase %", PCT, "=IF({al}>0,{inc}/{al},0)", "", ""),
        ("nb", "New base", EUR, "={al}+{inc}", "", ""),
        ("bmn", "Below range min after review", "0", "=IF({nb}<{bmin},1,0)", "Flag.", ""),
    ]),
    ("Step 6: Cost (annual, FTE-weighted)", [
        ("p0", "Total cash before equity, at actual FTE", EUR, "={tot}*{fte}", "(Annual base + bonus) × FTE.", ""),
        ("c0", "Employer contributions before equity", EUR, CONTRIB.replace("{x}", "{p0}"),
         "Full rate up to the contribution ceiling, tiered rates above it (Inputs).", ""),
        ("p1", "Total cash after equity, at actual FTE", EUR, "={al}*(1+{bonus})*{fte}", "", ""),
        ("c1", "Employer contributions after equity", EUR, CONTRIB.replace("{x}", "{p1}"), "", ""),
        ("p2", "Total cash after merit, at actual FTE", EUR, "={nb}*(1+{bonus})*{fte}", "", ""),
        ("c2", "Employer contributions after merit", EUR, CONTRIB.replace("{x}", "{p2}"), "", ""),
        ("s1b", "Equity: base", EUR, "={upb}*{fte}", "", ""),
        ("s1n", "Equity: bonus", EUR, "={upb}*{bonus}*{fte}", "Bonus grows with base.", ""),
        ("s1c", "Equity: employer contributions", EUR, "={c1}-{c0}", "Extra employer contributions on the increase.", ""),
        ("s2b", "Merit: base", EUR, "={inc}*{fte}", "", ""),
        ("s2n", "Merit: bonus", EUR, "={inc}*{bonus}*{fte}", "", ""),
        ("s2c", "Merit: employer contributions", EUR, "={c2}-{c1}", "", ""),
        ("eb", "Eligible base", EUR, "=IF({elig}=1,{al}*{fte},0)", "Aligned base of eligible staff, FTE-weighted: the denominator for the target %.", ""),
    ]),
    ("Step 6a: Target budget, the target % to every eligible person", [
        ("tinc", "Target increase", EUR, "=IF({elig}=1,{al}*PoolPct,0)", "Aligned base × target %, full time. A benchmark, not a payout.",
         ""),
        ("s3b", "Target: base", EUR, "={tinc}*{fte}", "", ""),
        ("s3n", "Target: bonus", EUR, "={tinc}*{bonus}*{fte}", "Each person's own bonus %.", ""),
        ("pT", "Total cash at target, at actual FTE", EUR, "=({al}+{tinc})*(1+{bonus})*{fte}", "", ""),
        ("cT", "Employer contributions at target", EUR, CONTRIB.replace("{x}", "{pT}"), "", ""),
        ("s3c", "Target: employer contributions", EUR, "={cT}-{c1}",
         "Depends on where each person sits against the ceiling, so it is not a flat rate on the target base.",
         ""),
    ]),
    ("Step 7: Total cash for the gap check", [
        ("ta1", "Total cash after equity, full time", EUR, "={al}*(1+{bonus})", "", ""),
        ("ta2", "Total cash after merit, full time", EUR, "={nb}*(1+{bonus})", "", ""),
        ("gy2", "grade×cash after merit", "0.00", "={grade}*{ta2}", "Helper sums for the regression after merit on Categories.", ""),
        ("ty2", "tenure×cash after merit", "0.00", "={ten}*{ta2}", "", ""),
    ]),
]

wc = wb.create_sheet("Calc")
col_of = {}
ci = 1
for step, items in CALC:
    for key, *_ in items:
        col_of[key] = get_column_letter(ci)
        ci += 1
DATA0 = 5

ci = 1
for step, items in CALC:
    first = True
    for key, header, fmt, tmpl, en, ru in items:
        if first:
            wc.cell(1, ci, step).font = BOLD
            first = False
        wc.cell(1, ci).fill = STEP
        wc.cell(2, ci, en).alignment = WRAP
        note2(wc, 3, ci, ru)
        wc.cell(4, ci, header).font = BOLD
        wc.cell(4, ci).alignment = WRAP
        for i in range(N):
            r = DATA0 + i
            refs = {k: f"{v}{r}" for k, v in col_of.items()}
            refs.update({"C_" + k: f"${v}:${v}" for k, v in col_of.items()})
            c = wc.cell(r, ci, tmpl.format(**refs))
            if fmt:
                c.number_format = fmt
        wc.column_dimensions[get_column_letter(ci)].width = 14
        ci += 1
wc.row_dimensions[2].height = 75
if any(item[5] for _, items in CALC for item in items):
    wc.row_dimensions[3].height = 75
wc.freeze_panes = "B5"
L = col_of["ovmax1"]
wc.conditional_formatting.add(f"{L}{DATA0}:{L}{DATA0 + N - 1}", CellIsRule(operator="equal", formula=["1"], fill=AMBER))


def CC(key):
    """Whole Calc column, for SUMIFS/COUNTIFS."""
    L = col_of[key]
    return f"Calc!${L}:${L}"


# ---------------------------------------------------------------- Categories
wk = wb.create_sheet("Categories")
cats = []
for r in rows:
    if r["category"] not in cats:
        cats.append(r["category"])
A = "$A{r}"
CAT_COLS = [
    ("Pay equity per category: unexplained gap closed in full", [
        ("name", "Category", None, None, "Type your categories here.", ""),
        ("n", "Headcount", "0", f"=COUNTIFS({CC('cat')},{A})", "", ""),
        ("nf", "Women", "0", f'=COUNTIFS({CC("cat")},{A},{CC("gen")},"F")', "", ""),
        ("nm", "Men", "0", f'=COUNTIFS({CC("cat")},{A},{CC("gen")},"M")', "", ""),
        ("mm", "Mean total cash, men", EUR, f'=IF({{nm}}>0,SUMIFS({CC("tot")},{CC("cat")},{A},{CC("gen")},"M")/{{nm}},0)', "Base + bonus, full time.", ""),
        ("mf", "Mean total cash, women", EUR, f'=IF({{nf}}>0,SUMIFS({CC("tot")},{CC("cat")},{A},{CC("gen")},"F")/{{nf}},0)', "", ""),
        ("raw", "Raw gap", PCT, "=IF(AND({nf}>0,{nm}>0,{mm}>0),({mm}-{mf})/{mm},0)", "(men − women) / men.", ""),
    ]),
    ("Regression of total cash on grade and tenure (explicit sums)", [
        ("sg", "Σ grade", "0.00", f"=SUMIFS({CC('grade')},{CC('cat')},{A})", "Sums over the category.", ""),
        ("st", "Σ tenure", "0.00", f"=SUMIFS({CC('ten')},{CC('cat')},{A})", "", ""),
        ("sy", "Σ total cash", "0.00", f"=SUMIFS({CC('tot')},{CC('cat')},{A})", "", ""),
        ("sgg", "Σ grade²", "0.00", f"=SUMIFS({CC('g2')},{CC('cat')},{A})", "", ""),
        ("stt", "Σ tenure²", "0.00", f"=SUMIFS({CC('t2')},{CC('cat')},{A})", "", ""),
        ("sgt", "Σ grade×tenure", "0.00", f"=SUMIFS({CC('gt')},{CC('cat')},{A})", "", ""),
        ("sgy", "Σ grade×cash", "0.00", f"=SUMIFS({CC('gy')},{CC('cat')},{A})", "", ""),
        ("sty", "Σ tenure×cash", "0.00", f"=SUMIFS({CC('ty')},{CC('cat')},{A})", "", ""),
        ("cgg", "Sgg", "0.0000", "=IF({n}>0,{sgg}-{sg}^2/{n},0)", "Centred sums.", ""),
        ("ctt", "Stt", "0.0000", "=IF({n}>0,{stt}-{st}^2/{n},0)", "", ""),
        ("cgt", "Sgt", "0.0000", "=IF({n}>0,{sgt}-{sg}*{st}/{n},0)", "", ""),
        ("cgy", "Sgy", "0.00", "=IF({n}>0,{sgy}-{sg}*{sy}/{n},0)", "", ""),
        ("cty", "Sty", "0.00", "=IF({n}>0,{sty}-{st}*{sy}/{n},0)", "", ""),
        ("det", "Determinant", "0.0000", "={cgg}*{ctt}-{cgt}^2", "Cramer's rule.", ""),
        ("ok", "Regression usable", "0", "=IF(AND({n}>=4,ABS({det})>1E-9),1,0)", "", ""),
        ("bg", "€ per grade", EUR, "=IF({ok}=1,({cgy}*{ctt}-{cty}*{cgt})/{det},0)", "", ""),
        ("bt", "€ per tenure year", EUR, "=IF({ok}=1,({cty}*{cgg}-{cgy}*{cgt})/{det},0)", "", ""),
    ]),
    ("Explained and unexplained gap", [
        ("gmm", "Mean grade, men", "0.00", f'=IF({{nm}}>0,SUMIFS({CC("grade")},{CC("cat")},{A},{CC("gen")},"M")/{{nm}},0)', "", ""),
        ("gmf", "Mean grade, women", "0.00", f'=IF({{nf}}>0,SUMIFS({CC("grade")},{CC("cat")},{A},{CC("gen")},"F")/{{nf}},0)', "", ""),
        ("tmm", "Mean tenure, men", "0.00", f'=IF({{nm}}>0,SUMIFS({CC("ten")},{CC("cat")},{A},{CC("gen")},"M")/{{nm}},0)', "", ""),
        ("tmf", "Mean tenure, women", "0.00", f'=IF({{nf}}>0,SUMIFS({CC("ten")},{CC("cat")},{A},{CC("gen")},"F")/{{nf}},0)', "", ""),
        ("exr", "Explained, before bounds", PCT, "=IF(AND({ok}=1,{nf}>0,{nm}>0,{mm}>0),({bg}*({gmm}-{gmf})+{bt}*({tmm}-{tmf}))/{mm},0)", "Part of the gap explained by differences in grade and tenure.", ""),
        ("ex", "Explained", PCT, "=IF({raw}>=0,MAX(0,MIN({exr},{raw})),MIN(0,MAX({exr},{raw})))", "Kept between 0 and the raw gap.", ""),
        ("un", "Unexplained", PCT, "={raw}-{ex}", "Closed in full when it is against women (positive). A gap in women's favour is not adjusted.", ""),
    ]),
    ("Who receives it and how much", [
        ("k1", "Median rank 1", "0", "=INT(({n}+1)/2)", "The middle place(s) in the category list sorted by total cash. Odd headcount: both ranks are the same person.",
         ""),
        ("k2", "Median rank 2", "0", "=INT({n}/2)+1", "", ""),
        ("med", "Median total cash", EUR, f"=(SUMIFS({CC('tot')},{CC('cat')},{A},{CC('rank')},{{k1}})+SUMIFS({CC('tot')},{CC('cat')},{A},{CC('rank')},{{k2}}))/2",
         "Category median of total cash: the average of the two middle values (one value twice if the headcount is odd).",
         ""),
        ("need", "Amount to close", EUR, "=IF(AND({raw}>=0,{un}>1E-12),{un}*{mm}*{nf},0)", "Unexplained gap × men's mean total cash × number of women.", ""),
        ("tdef", "Sum of shortfalls", EUR, f"=SUMIFS({CC('def')},{CC('cat')},{A})", "", ""),
    ]),
    ("Gap after each step", [
        ("m1", "Mean men, after equity", EUR, f'=IF({{nm}}>0,SUMIFS({CC("ta1")},{CC("cat")},{A},{CC("gen")},"M")/{{nm}},0)', "", ""),
        ("f1", "Mean women, after equity", EUR, f'=IF({{nf}}>0,SUMIFS({CC("ta1")},{CC("cat")},{A},{CC("gen")},"F")/{{nf}},0)', "", ""),
        ("gap1", "Raw gap after equity", PCT, "=IF({m1}>0,({m1}-{f1})/{m1},0)", "", ""),
        ("m2", "Mean men, after merit", EUR, f'=IF({{nm}}>0,SUMIFS({CC("ta2")},{CC("cat")},{A},{CC("gen")},"M")/{{nm}},0)', "", ""),
        ("f2", "Mean women, after merit", EUR, f'=IF({{nf}}>0,SUMIFS({CC("ta2")},{CC("cat")},{A},{CC("gen")},"F")/{{nf}},0)', "", ""),
        ("gap2", "Raw gap after merit", PCT, "=IF({m2}>0,({m2}-{f2})/{m2},0)", "Did the review widen the gap again?", ""),
    ]),
    ("Unexplained gap after each step", [
        ("un1", "Unexplained after equity", PCT, "=IF(AND({need}>0,{tdef}>0),0,{un})",
         "Zero where pay equity closed it; otherwise unchanged.", ""),
        ("sy2", "Σ total cash after merit", "0.00", f"=SUMIFS({CC('ta2')},{CC('cat')},{A})", "The same regression, re-run on total cash after merit.", ""),
        ("sgy2", "Σ grade×cash after merit", "0.00", f"=SUMIFS({CC('gy2')},{CC('cat')},{A})", "", ""),
        ("sty2", "Σ tenure×cash after merit", "0.00", f"=SUMIFS({CC('ty2')},{CC('cat')},{A})", "", ""),
        ("cgy2", "Sgy after merit", "0.00", "=IF({n}>0,{sgy2}-{sg}*{sy2}/{n},0)", "", ""),
        ("cty2", "Sty after merit", "0.00", "=IF({n}>0,{sty2}-{st}*{sy2}/{n},0)", "", ""),
        ("bg2", "€ per grade after merit", EUR, "=IF({ok}=1,({cgy2}*{ctt}-{cty2}*{cgt})/{det},0)", "", ""),
        ("bt2", "€ per tenure year after merit", EUR, "=IF({ok}=1,({cty2}*{cgg}-{cgy2}*{cgt})/{det},0)", "", ""),
        ("exr2", "Explained after merit, before bounds", PCT, "=IF(AND({ok}=1,{nf}>0,{nm}>0,{m2}>0),({bg2}*({gmm}-{gmf})+{bt2}*({tmm}-{tmf}))/{m2},0)", "", ""),
        ("ex2", "Explained after merit", PCT, "=IF({gap2}>=0,MAX(0,MIN({exr2},{gap2})),MIN(0,MAX({exr2},{gap2})))", "Kept between 0 and the raw gap.", ""),
        ("un2", "Unexplained after merit", PCT, "={gap2}-{ex2}", "Did the review open an unexplained gap again?", ""),
    ]),
]
kcol = {}
ci = 1
for _, items in CAT_COLS:
    for key, *_ in items:
        kcol[key] = get_column_letter(ci)
        ci += 1
# Calc refers to these columns by letter
assert (kcol["med"], kcol["need"], kcol["tdef"]) == ("AH", "AI", "AJ"), (kcol["med"], kcol["need"], kcol["tdef"])
ci = 1
for step, items in CAT_COLS:
    first = True
    for key, header, fmt, tmpl, en, ru in items:
        if first:
            wk.cell(1, ci, step).font = BOLD
            first = False
        wk.cell(1, ci).fill = STEP
        wk.cell(2, ci, en).alignment = WRAP
        note2(wk, 3, ci, ru)
        wk.cell(4, ci, header).font = BOLD
        wk.cell(4, ci).alignment = WRAP
        for i, cat in enumerate(cats):
            r = DATA0 + i
            if tmpl is None:
                c = wk.cell(r, ci, cat)
                c.fill = INPUT
            else:
                refs = {k: f"{v}{r}" for k, v in kcol.items()}
                refs["r"] = r
                c = wk.cell(r, ci, tmpl.replace("$A{r}", f"$A{r}").format(**refs))
            if fmt:
                c.number_format = fmt
        wk.column_dimensions[get_column_letter(ci)].width = 14
        ci += 1
wk.column_dimensions["A"].width = 22
wk.row_dimensions[2].height = 60
if any(item[5] for _, items in CAT_COLS for item in items):
    wk.row_dimensions[3].height = 60
wk.freeze_panes = "B5"

# ---------------------------------------------------------------- Matrix
wm = wb.create_sheet("Matrix")
wm["A1"] = "Merit matrix: the result, computed person by person"
wm["A1"].font = Font(bold=True, size=13)
wm["A2"] = "Rows: rating. Columns: compa-ratio (CR) zone after equity. Eligible staff only."
note2(wm, 3, 1, "", wrap=False)
blocks = [("Merit increase, % of eligible base payroll", PCT, "pct"), ("Headcount", "0", "n"),
          ("Merit increase, € (annual, FTE-weighted)", EUR, "eur")]
r0 = 5
for title, fmt, kind in blocks:
    wm.cell(r0, 1, title).font = BOLD
    for z in range(1, len(S["zones"]) + 1):
        nz = len(S["zones"])
        lo = f'TEXT(INDEX(Inputs!$B${Z1}:$B${Z2},{z}),"0%")'
        hi = f'TEXT(INDEX(Inputs!$B${Z1}:$B${Z2},{z + 1}),"0%")'
        head = f'="CR < "&{hi}' if z == 1 else f'="CR ≥ "&{lo}' if z == nz else f'="CR "&{lo}&" to "&{hi}'
        c = wm.cell(r0 + 1, 1 + z, head)
        c.font = BOLD
    for k in range(1, 6):
        r = r0 + 1 + k
        wm.cell(r, 1, f'=INDEX(Inputs!$A${R1}:$A${R2},{k})&" "&INDEX(Inputs!$B${R1}:$B${R2},{k})')
        for z in range(1, len(S["zones"]) + 1):
            rk = f"INDEX(Inputs!$A${R1}:$A${R2},{k})"
            crit = f'{CC("rating")},{rk},{CC("zone")},{z},{CC("elig")},1'
            if kind == "n":
                f = f"=COUNTIFS({crit})"
            elif kind == "eur":
                f = f"=SUMIFS({CC('s2b')},{crit})"
            else:
                f = f"=IF(SUMIFS({CC('eb')},{crit})>0,SUMIFS({CC('s2b')},{crit})/SUMIFS({CC('eb')},{crit}),0)"
            c = wm.cell(r, 1 + z, f)
            c.number_format = fmt
    r0 += 9
wm.column_dimensions["A"].width = 26
for col in "BCDE":
    wm.column_dimensions[col].width = 14

# ---------------------------------------------------------------- Summary
wsum = wb.create_sheet("Summary", 0)
wsum["A1"] = "Summary: what the review costs"
wsum["A1"].font = Font(bold=True, size=13)
wsum.append([])
wsum.append(["", "Value"])
note2(wsum, 3, 3, "", wrap=False)
for c in wsum[3]:
    c.font = BOLD


def S_(key):
    return f"SUM({CC(key)})"


# (key, label, formula, number format, second-language note); {key} -> that line's value cell
lines = [
    ("hc", "Headcount", f"=COUNT({CC('base')})", "0", ""),
    ("ne", "Eligible for merit", f"={S_('elig')}", "0", ""),
    (None,),
    (None, "PAYROLL COST: full year, incl. employer contributions", None, None, ""),
    ("pr0", "Current payroll cost", f"={S_('p0')}+{S_('c0')}", EUR, ""),
    ("pay0", "   of which total cash (base + bonus)", f"={S_('p0')}", EUR, ""),
    ("con0", "   of which employer contributions", f"={S_('c0')}", EUR, ""),
    ("pr1", "Payroll cost after pay equity", f"={S_('p1')}+{S_('c1')}", EUR, ""),
    ("pr2", "Payroll cost after merit review", f"={S_('p2')}+{S_('c2')}", EUR, ""),
    ("d1", "Change: pay equity", "={pr1}-{pr0}", EUR, ""),
    ("d2", "Change: merit review", "={pr2}-{pr1}", EUR, ""),
    ("dt", "Change: total", "={pr2}-{pr0}", EUR, ""),
    ("dtp", "Change: total, % of current payroll cost", "=IF({pr0}>0,{dt}/{pr0},0)", PCT, ""),
    (None,),
    (None, "THIS YEAR: from the effective month", None, None, ""),
    ("t1", "Change this year: pay equity", "={d1}*InYearFactor", EUR, ""),
    ("t2", "Change this year: merit review", "={d2}*InYearFactor", EUR, ""),
    ("tt", "Change this year: total", "={t1}+{t2}", EUR, ""),
    ("ttp", "Change this year, % of current payroll cost", "=IF({pr0}>0,{tt}/{pr0},0)", PCT, ""),
    (None,),
    (None, "PAY EQUITY ADJUSTMENT (Calc step 2), breakdown", None, None, ""),
    ("a_b", "Base adjustment", f"={S_('s1b')}", EUR, ""),
    ("a_n", "Bonus effect", f"={S_('s1n')}", EUR, ""),
    ("a_c", "Employer contributions", f"={S_('s1c')}", EUR, ""),
    ("a_y", "Total, full year", "={a_b}+{a_n}+{a_c}", EUR, ""),
    ("a_t", "Total, this year", "={a_y}*InYearFactor", EUR, ""),
    (None,),
    (None, "MERIT INCREASE from the aligned base (Calc step 5), breakdown", None, None, ""),
    ("b_b", "Merit increase (base)", f"={S_('s2b')}", EUR, ""),
    ("b_n", "Bonus effect", f"={S_('s2n')}", EUR, ""),
    ("b_c", "Employer contributions", f"={S_('s2c')}", EUR, ""),
    ("b_y", "Total, full year", "={b_b}+{b_n}+{b_c}", EUR, ""),
    ("b_t", "Total, this year", "={b_y}*InYearFactor", EUR, ""),
    (None,),
    (None, "MERIT REVIEW VS TARGET BUDGET, incl. bonus and employer contributions", None, None,
     ""),
    ("eb", "Eligible base payroll", f"={S_('eb')}", EUR, ""),
    ("pp", "Target merit increase, % of eligible base payroll (change on Inputs)", "=PoolPct", PCT,
     ""),
    ("mp", "Merit increase as calculated, % of eligible base payroll", "=IF({eb}>0,{b_b}/{eb},0)", PCT, ""),
    ("dpp", "Variance vs target, pp (+ over, − under)", "=({mp}-{pp})*100", "0.00", ""),
    ("g_y", "Target budget, full year", "={g_b}+{g_n}+{g_c}", EUR, ""),
    ("g_b", "   of which base", f"={S_('s3b')}", EUR, ""),
    ("g_n", "   of which bonus", f"={S_('s3n')}", EUR, ""),
    ("g_c", "   of which employer contributions", f"={S_('s3c')}", EUR, ""),
    ("deur", "Variance vs target budget, € full year (+ over)", "={b_y}-{g_y}", EUR, ""),
    ("g_t", "Target budget, this year", "={g_y}*InYearFactor", EUR, ""),
    ("deut", "Variance vs target budget, € this year (+ over)", "={b_t}-{g_t}", EUR, ""),
    (None,),
    (None, "FLAGS", None, None, ""),
    ("f1", "Above range max after equity adjustment (amber: check, not an error)", f"={S_('ovmax1')}", "0", ""),
    ("f2", "At or above range max: no merit increase (red-circled if above)", f"={S_('red')}", "0", ""),
    ("f3", "Below range min after review", f"={S_('bmn')}", "0", ""),
    (None,),
    (None, "DATA CHECKS: all should be 0", None, None, ""),
    ("k1", "Employee rows not in the calculation", f"=COUNTA(Employees!$A:$A)-1-{{hc}}", "0",
     ""),
    ("k2", "Employees without a range on Bands", f"={{hc}}-{S_('bok')}", "0",
     ""),
    ("k3", "Employees in categories missing from Categories", f"={{hc}}-SUM(Categories!$B:$B)", "0",
     ""),
    ("k4", "Employees with a rating missing from the rating table", f"={{hc}}-SUMPRODUCT(COUNTIFS({CC('rating')},Inputs!$A${R1}:$A${R2}))", "0",
     ""),
]
srow = {ln[0]: 4 + i for i, ln in enumerate(lines) if len(ln) > 1 and ln[0]}
for i, ln in enumerate(lines):
    r = 4 + i
    if len(ln) == 1:
        continue
    key, label, f, fmt, ru = ln
    wsum.cell(r, 1, label)
    if f is None:
        wsum.cell(r, 1).font = BOLD
        note2(wsum, r, 3, ru, Font(bold=True, italic=True, color="7F7F7F"), wrap=False)
        continue
    c = wsum.cell(r, 2, f.format(**{k: f"B{v}" for k, v in srow.items()}))
    c.number_format = fmt
    note2(wsum, r, 3, ru, wrap=r == srow.get("f1"))
wsum.column_dimensions["A"].width = 46
wsum.column_dimensions["B"].width = 18
wsum.column_dimensions["C"].width = 48
c = wsum.cell(srow["pp"], 1)
c.hyperlink = "#Inputs!B%d" % (4 + [x[0] for x in settings].index("PoolPct"))
c.font = Font(color="0563C1", underline="single")
for k in ("k1", "k2", "k3", "k4"):
    wsum.conditional_formatting.add(f"B{srow[k]}", CellIsRule(operator="notEqual", formula=["0"], fill=AMBER))
r = srow["f1"]
wsum.conditional_formatting.add(f"B{r}", CellIsRule(operator="greaterThan", formula=["0"], fill=AMBER))

wb.save(OUT)
print("written", OUT, "rows", N, "categories", len(cats))
