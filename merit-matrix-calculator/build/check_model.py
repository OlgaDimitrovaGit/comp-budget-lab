# -*- coding: utf-8 -*-
"""
Checks the recalculated workbook against reference_calc.py.

Run after Excel has recalculated merit-model.xlsx. Exits 1 on any formula
error or any mismatch above one cent.
"""
import sys

from openpyxl import load_workbook

from reference_calc import MODEL, SAMPLE_EMPLOYEES, SAMPLE_RANGES, analyse, load, load_bands

SRC = sys.argv[1] if len(sys.argv) > 1 else SAMPLE_EMPLOYEES
BANDS = sys.argv[2] if len(sys.argv) > 2 else SAMPLE_RANGES
XLSX = sys.argv[3] if len(sys.argv) > 3 else MODEL

wb = load_workbook(XLSX, data_only=True)
fail = 0

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("#"):
                print("ERROR %s!%s %s" % (ws.title, c.coordinate, c.value))
                fail += 1

ref = analyse(load(SRC), load_bands(BANDS))
s = wb["Summary"]
rows = {}
for r in range(1, s.max_row + 1):
    lab = s.cell(r, 1).value
    if lab:
        rows.setdefault(lab, []).append(r)


def B(label, nth=0):
    """Value next to a Summary label; nth picks a repeated label (step A = 0, step B = 1)."""
    return s.cell(rows[label][nth], 2).value


expect = [
    (("Current payroll cost",), ref["payroll"], "starting budget"),
    (("   of which total cash (base + bonus)",), ref["pay_before"], "of which pay"),
    (("   of which employer contributions",), ref["contrib_before"], "of which contributions"),
    (("Payroll cost after pay equity",), ref["payroll"] + ref["step1"]["annual"], "after equity"),
    (("Payroll cost after merit review",), ref["payroll"] + ref["cost_annual"], "after merit"),
    (("Change: pay equity",), ref["step1"]["annual"], "change equity"),
    (("Change: merit review",), ref["step2"]["annual"], "change merit"),
    (("Change: total",), ref["cost_annual"], "change total"),
    (("Change: total, % of current payroll cost",), ref["load_full_pct"] / 100, "change total %"),
    (("Change this year: total",), ref["cost_in_year"], "change this year"),
    (("Change this year, % of current payroll cost",), ref["load_in_year_pct"] / 100, "change this year %"),
    (("Base adjustment",), ref["step1"]["base"], "equity base"),
    (("Bonus effect", 0), ref["step1"]["bonus"], "equity bonus"),
    (("Employer contributions", 0), ref["step1"]["contrib"], "equity contributions"),
    (("Total, full year", 0), ref["step1"]["annual"], "equity full year"),
    (("Total, this year", 0), ref["step1"]["in_year"], "equity this year"),
    (("Merit increase (base)",), ref["step2"]["base"], "merit base"),
    (("Bonus effect", 1), ref["step2"]["bonus"], "merit bonus"),
    (("Employer contributions", 1), ref["step2"]["contrib"], "merit contributions"),
    (("Total, full year", 1), ref["step2"]["annual"], "merit full year"),
    (("Total, this year", 1), ref["step2"]["in_year"], "merit this year"),
    (("Eligible base payroll",), ref["eligible_base"], "eligible base"),
    (("   of which base",), ref["target"]["base"], "target base"),
    (("   of which bonus",), ref["target"]["bonus"], "target bonus"),
    (("   of which employer contributions", 1), ref["target"]["contrib"], "target contributions"),
    (("Target budget, full year",), ref["target"]["annual"], "target full year"),
    (("Target budget, this year",), ref["target"]["in_year"], "target this year"),
    (("Variance vs target, pp (+ over, − under)",), ref["dev_pp"], "deviation pp"),
    (("Variance vs target budget, € full year (+ over)",), ref["dev_eur"], "deviation EUR"),
    (("Variance vs target budget, € this year (+ over)",), ref["dev_eur_in_year"], "deviation EUR this year"),
    (("Above range max after equity adjustment (amber: check, not an error)",), ref["flag_equity_over_max"], "flag equity>max"),
    (("Below range min after review",), ref["below_min_after"], "below min after review"),
    (("At or above range max: no merit increase (red-circled if above)",), ref["red_circled"], "red-circled"),
]
for key, want, label in expect:
    got = B(*key)
    tol = 1e-6 if abs(want) < 1 else 0.01  # shares vs euro amounts
    ok = got is not None and abs(got - want) < tol
    print("%-4s %-22s excel=%14.4f ref=%14.4f" % ("OK" if ok else "FAIL", label, got or 0, want))
    fail += 0 if ok else 1

c = wb["Calc"]
hdr = {c.cell(4, j).value: j for j in range(1, c.max_column + 1)}
for i, e in enumerate(ref["emp"]):
    r = 5 + i
    for h, key in (("Aligned base", "aligned"), ("Merit increase", "increase"), ("Compa-ratio zone", "zone"),
                   ("Eligible", "eligible")):
        got, want = c.cell(r, hdr[h]).value, float(e[key])
        if got is None or abs(got - want) > 0.01:
            print("FAIL %s %s excel=%s ref=%s" % (e["id"], h, got, want))
            fail += 1

k = wb["Categories"]
kh = {k.cell(4, j).value: j for j in range(1, k.max_column + 1)}
krow = {k.cell(r, 1).value: r for r in range(5, k.max_row + 1) if k.cell(r, 1).value}
for name, cat in ref["cats"].items():
    for h, key in (("Raw gap", "raw"), ("Unexplained", "unexplained"),
                   ("Unexplained after equity", "unexplained_after_equity"),
                   ("Unexplained after merit", "unexplained_after_merit")):
        got = k.cell(krow[name], kh[h]).value if name in krow else None
        if got is None or abs(got * 100 - cat[key]) > 1e-6:
            print("FAIL %s %s excel=%s ref=%.6f" % (name, h, got, cat[key]))
            fail += 1

print("FAILURES:", fail)
sys.exit(1 if fail else 0)
